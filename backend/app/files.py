from io import BytesIO
from pathlib import Path
from fastapi import HTTPException, UploadFile
from pypdf import PdfReader
from docx import Document
from .config import settings

async def extract_transcript(upload: UploadFile) -> str:
    name = Path(upload.filename or "").name
    ext = Path(name).suffix.lower()
    if ext not in {".txt", ".pdf", ".docx"}:
        raise HTTPException(415, "Unsupported file type. Upload a TXT, PDF, or DOCX document.")
    max_bytes = settings.transcript_max_size_mb * 1024 * 1024
    content = await upload.read(max_bytes + 1)
    if len(content) > max_bytes:
        raise HTTPException(413, f"File exceeds the {settings.transcript_max_size_mb} MB limit.")
    if not content:
        raise HTTPException(422, "The uploaded document is empty.")
    try:
        if ext == ".txt":
            text = content.decode("utf-8-sig")
        elif ext == ".pdf":
            reader = PdfReader(BytesIO(content))
            text = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
            if not text:
                raise HTTPException(422, "This PDF has no selectable text. Scanned PDFs need OCR before upload.")
        else:
            doc = Document(BytesIO(content))
            text = "\n".join(p.text for p in doc.paragraphs).strip()
            text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, "We could not read this document. Check that it is a valid, unprotected TXT, PDF, or DOCX file.") from exc
    text = text.strip()
    if not text:
        raise HTTPException(422, "No transcript text was found in the uploaded document.")
    return text
